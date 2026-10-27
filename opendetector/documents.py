"""Text extraction and section replacement for Markdown, Word, and PowerPoint files."""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO


@dataclass(frozen=True)
class DocumentSection:
    label: str
    text: str
    kind: str = ""
    locator: tuple[int, ...] = ()


@dataclass(frozen=True)
class AnswerSlot:
    label: str
    question: str
    previous_answer: str
    target: DocumentSection


def extract_document(filename: str, data: bytes) -> list[DocumentSection]:
    """Extract non-empty document sections while retaining their source locations."""
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "md":
        try:
            return _markdown_sections(data)
        except UnicodeDecodeError as exc:
            raise ValueError("The Markdown file must use UTF-8 encoding.") from exc
    readers = {"docx": _docx_sections, "pptx": _pptx_sections}
    if suffix not in readers:
        raise ValueError("Upload a .md Markdown file, .docx Word document, or .pptx PowerPoint presentation.")
    try:
        return readers[suffix](BytesIO(data))
    except ImportError:
        raise
    except Exception as exc:  # damaged or renamed files raise zip, XML and package errors
        raise ValueError(f"It isn't a valid .{suffix} file. It may be damaged, or another format "
                         f"renamed to .{suffix}.") from exc


def extract_answer_slots(filename: str, data: bytes) -> list[AnswerSlot]:
    """Find question/answer table rows and side-by-side PowerPoint text boxes.

    The leftmost table cell is treated as the question and the rightmost as its answer box.
    PowerPoint text boxes are paired by left/right position and vertical alignment.
    """
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "md":
        try:
            lines = data.decode("utf-8-sig").splitlines(keepends=True)
        except UnicodeDecodeError as exc:
            raise ValueError("The Markdown file must use UTF-8 encoding.") from exc
        slots = []
        table_number = 0
        row_number = 0
        in_table = False
        for line_number, line in enumerate(lines):
            body = line.rstrip("\r\n")
            cells = _markdown_cell_spans(body)
            if len(cells) < 2:
                in_table = False
                continue
            if not in_table:
                table_number += 1
                row_number = 0
                in_table = True
            if _is_markdown_separator([cell[2] for cell in cells]):
                continue
            row_number += 1
            question, answer = cells[0][2].strip(), cells[-1][2].strip()
            if not question or _is_table_header(question, answer):
                continue
            target = DocumentSection(
                f"Table {table_number}, row {row_number}, answer cell",
                answer,
                "markdown_table_cell",
                (line_number, cells[-1][0], cells[-1][1]),
            )
            slots.append(AnswerSlot(f"Table {table_number}, row {row_number}", question, answer, target))
        return slots

    if suffix == "docx":
        from docx import Document

        document = Document(BytesIO(data))
        slots = []
        for table_number, table in enumerate(document.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                if len(row.cells) < 2:
                    continue
                question, answer = row.cells[0].text.strip(), row.cells[-1].text.strip()
                if not question or _is_table_header(question, answer):
                    continue
                target = DocumentSection(
                    f"Table {table_number}, row {row_number}, cell {len(row.cells)}",
                    answer,
                    "docx_table_cell",
                    (table_number - 1, row_number - 1, len(row.cells) - 1),
                )
                slots.append(AnswerSlot(f"Table {table_number}, row {row_number}", question, answer, target))
        return slots

    if suffix == "pptx":
        from pptx import Presentation

        presentation = Presentation(BytesIO(data))
        slots = []
        for slide_number, slide in enumerate(presentation.slides, 1):
            for shape_number, shape in enumerate(slide.shapes, 1):
                if not getattr(shape, "has_table", False) or len(shape.table.columns) < 2:
                    continue
                for row_number, row in enumerate(shape.table.rows, 1):
                    question, answer = row.cells[0].text.strip(), row.cells[-1].text.strip()
                    if not question or _is_table_header(question, answer):
                        continue
                    target = DocumentSection(
                        f"Slide {slide_number}, Table row {row_number}, cell {len(row.cells)}",
                        answer,
                        "pptx_table_cell",
                        (slide_number - 1, shape_number - 1, row_number - 1, len(row.cells) - 1),
                    )
                    slots.append(AnswerSlot(f"Slide {slide_number}, table row {row_number}",
                                            question, answer, target))

            midpoint = presentation.slide_width // 2
            left_shapes, right_shapes = [], []
            for shape_number, shape in enumerate(slide.shapes, 1):
                if getattr(shape, "has_table", False) or not getattr(shape, "has_text_frame", False):
                    continue
                if shape.left < midpoint:
                    question = shape.text.strip()
                    if question:
                        left_shapes.append((shape_number, shape, question))
                else:
                    right_shapes.append((shape_number, shape, shape.text.strip()))
            used_right = set()
            for question_number, question_shape, question in left_shapes:
                q_center = question_shape.top + question_shape.height / 2
                candidates = []
                for shape_number, answer_shape, answer in right_shapes:
                    if shape_number in used_right:
                        continue
                    a_center = answer_shape.top + answer_shape.height / 2
                    tolerance = max(presentation.slide_height // 12,
                                    max(question_shape.height, answer_shape.height) // 2)
                    if abs(q_center - a_center) <= tolerance:
                        candidates.append((abs(q_center - a_center), shape_number, answer_shape, answer))
                if not candidates:
                    continue
                _, shape_number, answer_shape, answer = min(candidates, key=lambda item: item[0])
                used_right.add(shape_number)
                question_name = getattr(question_shape, "name", "") or f"Text box {question_number}"
                target_name = getattr(answer_shape, "name", "") or f"Text box {shape_number}"
                target = DocumentSection(
                    f"Slide {slide_number}, {target_name}", answer, "pptx_shape",
                    (slide_number - 1, shape_number - 1),
                )
                slots.append(AnswerSlot(f"Slide {slide_number}, {question_name}",
                                        question, answer, target))
        return slots

    raise ValueError("Answer box detection supports .md, .docx, and .pptx files.")


def _markdown_cell_spans(line: str) -> list[tuple[int, int, str]]:
    """Return cell content spans for a pipe-delimited Markdown table row."""
    pipe_positions = [match.start() for match in re.finditer(r"(?<!\\)\|", line)]
    if not pipe_positions:
        return []
    spans = []
    start = 0
    for pipe_index, end in enumerate(pipe_positions):
        segment = line[start:end]
        left_padding = len(segment) - len(segment.lstrip())
        cell_start = start + left_padding
        cell_end = start + len(segment) - (len(segment) - len(segment.rstrip()))
        if not segment.strip():
            cell_end = cell_start
        outer_leading_pipe = pipe_index == 0 and not segment.strip()
        if not outer_leading_pipe:
            spans.append((cell_start, cell_end, line[cell_start:cell_end]))
        start = end + 1
    segment = line[start:]
    if segment.strip():
        left_padding = len(segment) - len(segment.lstrip())
        cell_start = start + left_padding
        cell_end = start + len(segment) - (len(segment) - len(segment.rstrip()))
        spans.append((cell_start, cell_end, line[cell_start:cell_end]))
    return spans


def _is_markdown_separator(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def _markdown_sections(data: bytes) -> list[DocumentSection]:
    text = data.decode("utf-8-sig")
    lines = text.splitlines(keepends=True)
    sections = []
    index = 0
    while index < len(lines):
        body = lines[index].strip()
        if not body:
            index += 1
            continue
        cells = _markdown_cell_spans(lines[index].rstrip("\r\n"))
        if len(cells) >= 2:
            if _is_markdown_separator([cell[2] for cell in cells]):
                index += 1
                continue
            for cell_number, (start, end, value) in enumerate(cells, 1):
                if value:
                    sections.append(DocumentSection(
                        f"Table row at line {index + 1}, cell {cell_number}", value,
                        "markdown_table_cell", (index, start, end),
                    ))
            index += 1
            continue
        heading = re.match(r"^(\s{0,3}#{1,6}\s+)(.*?)(\s*#*\s*)$", lines[index].rstrip("\r\n"))
        if heading:
            sections.append(DocumentSection(
                f"Heading at line {index + 1}", heading.group(2).strip(),
                "markdown_heading", (index,),
            ))
            index += 1
            continue
        start = index
        while index < len(lines):
            current = lines[index].strip()
            if not current or len(_markdown_cell_spans(lines[index].rstrip("\r\n"))) >= 2:
                break
            if re.match(r"^\s{0,3}#{1,6}\s+", lines[index]):
                break
            index += 1
        section_text = "\n".join(line.strip() for line in lines[start:index]).strip()
        if section_text:
            sections.append(DocumentSection(
                f"Text block at line {start + 1}", section_text, "markdown_block", (start, index),
            ))
    return sections


def _is_table_header(question: str, answer: str) -> bool:
    q = question.strip().casefold()
    a = answer.strip().casefold()
    return q in {"question", "prompt", "question/prompt"} and a in {
        "", "answer", "response", "your answer"
    }


def _docx_sections(stream: BytesIO) -> list[DocumentSection]:
    from docx import Document

    document = Document(stream)
    sections = [
        DocumentSection(f"Paragraph {i}", paragraph.text.strip(), "docx_paragraph", (i - 1,))
        for i, paragraph in enumerate(document.paragraphs, 1)
        if paragraph.text.strip()
    ]
    for table_number, table in enumerate(document.tables, 1):
        for row_number, row in enumerate(table.rows, 1):
            for cell_number, cell in enumerate(row.cells, 1):
                value = cell.text.strip()
                if value:
                    sections.append(DocumentSection(
                        f"Table {table_number}, row {row_number}, cell {cell_number}",
                        value,
                        "docx_table_cell",
                        (table_number - 1, row_number - 1, cell_number - 1),
                    ))
    return sections


def _pptx_sections(stream: BytesIO) -> list[DocumentSection]:
    from pptx import Presentation

    presentation = Presentation(stream)
    sections = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        table_number = 0
        for shape_number, shape in enumerate(slide.shapes, 1):
            if getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value:
                    shape_name = getattr(shape, "name", "") or f"Text box {shape_number}"
                    sections.append(DocumentSection(
                        f"Slide {slide_number}, {shape_name}", value, "pptx_shape",
                        (slide_number - 1, shape_number - 1),
                    ))
            if getattr(shape, "has_table", False):
                table_number += 1
                for row_number, row in enumerate(shape.table.rows, 1):
                    for cell_number, cell in enumerate(row.cells, 1):
                        value = cell.text.strip()
                        if value:
                            sections.append(DocumentSection(
                                f"Slide {slide_number}, Table {table_number}, row {row_number}, cell {cell_number}",
                                value,
                                "pptx_table_cell",
                                (slide_number - 1, shape_number - 1, row_number - 1, cell_number - 1),
                            ))
    return sections


def replace_document_section(filename: str, data: bytes, section: DocumentSection, text: str) -> bytes:
    """Return a copy of a supported document with one extracted text section replaced in place."""
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == "md":
        had_bom = data.startswith(b"\xef\xbb\xbf")
        lines = data.decode("utf-8-sig").splitlines(keepends=True)
        if section.kind == "markdown_table_cell":
            line_index, start, end = section.locator
            line = lines[line_index]
            body = line.rstrip("\r\n")
            newline = line[len(body):]
            replacement = text.replace("|", r"\|").replace("\r", " ").replace("\n", "<br>")
            lines[line_index] = body[:start] + replacement + body[end:] + newline
        elif section.kind == "markdown_heading":
            (line_index,) = section.locator
            line = lines[line_index]
            body = line.rstrip("\r\n")
            newline = line[len(body):]
            heading = re.match(r"^(\s{0,3}#{1,6}\s+)(.*?)(\s*#*\s*)$", body)
            if not heading:
                raise ValueError("The selected Markdown heading could not be found in the source file.")
            lines[line_index] = heading.group(1) + text.replace("\n", " ") + heading.group(3) + newline
        elif section.kind == "markdown_block":
            start, end = section.locator
            old_lines = lines[start:end]
            newline = "\r\n" if any(line.endswith("\r\n") for line in old_lines) else "\n"
            trailing_newline = bool(old_lines and old_lines[-1].endswith(("\n", "\r")))
            replacement = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
            rendered = newline.join(replacement)
            if trailing_newline:
                rendered += newline
            lines[start:end] = [rendered]
        else:
            raise ValueError("The selected section cannot be replaced in this Markdown file.")
        encoded = "".join(lines).encode("utf-8")
        return (b"\xef\xbb\xbf" + encoded) if had_bom else encoded

    if suffix == "docx":
        from docx import Document

        document = Document(BytesIO(data))
        if section.kind == "docx_paragraph":
            paragraph = document.paragraphs[section.locator[0]]
            _replace_docx_paragraph(paragraph, text)
        elif section.kind == "docx_table_cell":
            table_index, row_index, cell_index = section.locator
            cell = document.tables[table_index].rows[row_index].cells[cell_index]
            _replace_docx_paragraph(cell.paragraphs[0], text)
            for paragraph in cell.paragraphs[1:]:
                _replace_docx_paragraph(paragraph, "")
        else:
            raise ValueError("The selected section cannot be replaced in this Word document.")
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    if suffix == "pptx":
        from pptx import Presentation

        presentation = Presentation(BytesIO(data))
        if section.kind == "pptx_shape":
            slide_index, shape_index = section.locator
            shape = presentation.slides[slide_index].shapes[shape_index]
            if not getattr(shape, "has_text_frame", False):
                raise ValueError("The selected PowerPoint shape no longer contains text.")
            _replace_pptx_text_frame(shape.text_frame, text)
        elif section.kind == "pptx_table_cell":
            slide_index, shape_index, row_index, cell_index = section.locator
            shape = presentation.slides[slide_index].shapes[shape_index]
            _replace_pptx_text_frame(shape.table.cell(row_index, cell_index).text_frame, text)
        else:
            raise ValueError("The selected section cannot be replaced in this PowerPoint file.")
        output = BytesIO()
        presentation.save(output)
        return output.getvalue()

    raise ValueError("Export supports .md, .docx Word documents, and .pptx PowerPoint files.")


def _replace_docx_paragraph(paragraph, text: str) -> None:
    """Keep paragraph and first-run formatting while replacing only its characters."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _replace_pptx_text_frame(text_frame, text: str) -> None:
    """Keep the selected text box and first-run styling, clearing only its old text."""
    paragraphs = text_frame.paragraphs
    if not paragraphs:
        text_frame.text = text
        return
    first = paragraphs[0]
    if first.runs:
        first.runs[0].text = text.replace("\n", " ")
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.add_run().text = text.replace("\n", " ")
    for paragraph in paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def _row_text(cells) -> str:
    """A table row as "a | b | c", or "" when every cell is blank."""
    values = [cell.text.strip() for cell in cells]
    return " | ".join(values) if any(values) else ""

"""Command-line tests with a stand-in detector, so no model downloads are needed."""

import io
import json

import pytest
from docx import Document

from opendetector import Detector, DetectorConfig
from opendetector.__main__ import main
from tests.test_detector import AI, HUMAN, KeywordClassifier, LengthBinoculars


@pytest.fixture
def stub_detector(monkeypatch):
    configs = []

    def from_config(cls, config=None, load_combiner=True):
        configs.append(config)
        return Detector([KeywordClassifier()], LengthBinoculars(), None, DetectorConfig(min_words=40))

    monkeypatch.setattr(Detector, "from_config", classmethod(from_config))
    return configs


@pytest.fixture
def worksheet(tmp_path):
    document = Document()
    document.add_paragraph("Instructions. " + HUMAN * 6)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Explain your plan.", AI * 6
    table.cell(1, 0).text, table.cell(1, 1).text = "Name", "Sam"
    path = tmp_path / "worksheet.docx"
    document.save(path)
    return path


def run(capsys, *argv):
    main(list(argv))
    return json.loads(capsys.readouterr().out)


def test_document_gets_a_verdict_and_long_sections_their_own(stub_detector, worksheet, capsys):
    report = run(capsys, str(worksheet), "--no-sentences")
    assert report["ai_probability"] is not None and report["words"] > 150
    assert report["sections_extracted"] == 5
    assert [s["label"] for s in report["sections"]] == ["Paragraph 1", "Table 1, row 1, cell 2"]
    assert stub_detector[0].glyph_model is not None  # full profile by default


def test_answers_only_leaves_out_the_template(stub_detector, worksheet, capsys):
    report = run(capsys, str(worksheet), "--answers-only", "--profile", "slim", "--device", "mps")
    assert report["sections_extracted"] == 2  # the long answer and "Sam"
    assert all("grandmother" not in s["text"] for s in report["sentences"])
    assert stub_detector[0].glyph_model is None and stub_detector[0].device == "mps"


def test_answers_only_rejects_plain_text_and_empty_worksheets(stub_detector, tmp_path):
    essay = tmp_path / "essay.txt"
    essay.write_text(HUMAN * 6)
    with pytest.raises(SystemExit):
        main([str(essay), "--answers-only"])
    blank = tmp_path / "blank.docx"
    Document().save(blank)
    with pytest.raises(SystemExit):
        main([str(blank), "--answers-only"])
    assert stub_detector == []


def test_plain_text_from_stdin_has_no_sections(stub_detector, monkeypatch, capsys):
    monkeypatch.setattr("sys.stdin", io.StringIO(HUMAN * 6))
    report = run(capsys)
    assert report["verdict"] == "likely_human" and "sections" not in report

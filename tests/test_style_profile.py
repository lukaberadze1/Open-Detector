import json

import pytest

from update_style_profile import (
    FLOOR,
    banned_patterns,
    finished_pieces,
    format_errors,
    main,
    measure,
    parse_note,
    register_numbers,
    summarize,
    week_weight,
)

NOTE = """---
id: "202608291209"
created: "29-08-2026"
---
# Why Rome?
> [!info]- 🧭 Split out of [[Roman Republic]] on 13 Oct 2026, verbatim.
the review keeps going on this line without its marker
<!-- a template prompt the writer never filled in -->

Rome is a wild city, normally cities grow from a good position such as [[Trade Routes|trade routes]] ^wild-city
![[Screenshot 2026-08-29.png]]
==Archimedes, then in his mid-70s, had made Syracuse nearly impregnable.==
`Alexander the Great cut the **Gordian Knot**, a legendary knot
`- **The Act**: In **333 BC**, Alexander famously sliced it in half
Numitor’s daughter gave birth to twin sons
People — [[Alexander the Great]] · [[Pericles]]
*Book's question: why did Rome live in near-constant war?*
- **Who paid / who was bound by it:**
- [ ] process this capture
- Sparta had two kings
- the phalanx was not formed from mercanaries but from free citizens
ახალი წელი დაიწყო ძალიან კარგად
Sparta had two kings.
"""


def test_parse_note_keeps_only_the_writers_own_lines():
    week, paragraphs, pasted = parse_note(NOTE)
    assert week == "2026-W35"
    assert paragraphs == [
        "Rome is a wild city, normally cities grow from a good position such as trade routes",
        "the phalanx was not formed from mercanaries but from free citizens",
        "Sparta had two kings.",
    ]
    assert len(pasted) == 4  # two backtick lines, the curly quote, the dashed link list
    assert all("Archimedes" not in line for line in pasted)  # a highlight is dropped outright


def test_parse_note_reads_iso_dates_and_notes_without_frontmatter():
    line = "The tithe was far older than Heiro II."
    assert parse_note(f'---\ncreated: "2026-10-14 08:10"\n---\n{line}\n')[0] == "2026-W42"
    assert parse_note(line) == (None, [line], [])


def test_week_weight_halves_and_never_drops_below_the_floor():
    assert week_weight("2026-W44", "2026-W44") == 1.0
    assert week_weight("2026-W32", "2026-W44") == pytest.approx(0.5)
    assert week_weight("2024-W01", "2026-W44") == FLOOR
    assert week_weight("2026-W45", "2026-W44") == 1.0  # a future week counts as this week


def test_measure_counts_the_writers_fingerprint():
    texts = ["However this is evident. However, not here at all.",
             "It was a ideal battleground and its clear that he dont care.",
             "Even so we delve into it."]
    counts = measure(texts, {"delve": __import__("re").compile(r"\bdelve\b", 2)})
    s = summarize(counts)
    assert s["mechanics"]["however_followed_by_comma"] == "1 of 2"
    assert s["mechanics"]["a_vs_an"] == "1 a, 0 an"
    assert s["mechanics"]["its_for_it_is_vs_it_s"] == "1 its, 0 it's"
    assert s["mechanics"]["contractions_without_apostrophe_pct"] == 100.0
    assert s["sentence_length"]["min"] == 4 and s["sentence_length"]["max"] == 12
    assert s["banned_phrases_used"] == {"delve": 1}


def test_registers_keep_notes_and_finished_work_apart():
    alone = register_numbers({})
    assert alone["finalized"]["avg_sentence_length_words"] == 18.3 and "notes" not in alone
    long_notes = measure([" ".join(["word"] * 40) + "."] * 50, {})
    numbers = register_numbers(long_notes)
    assert numbers["notes"]["avg_sentence_length_words"] == 40.0, "notes are measured on their own, not blended"
    assert numbers["notes"]["sentence_length_distribution"]["30_plus_words_pct"] == "100.0%"
    assert numbers["finalized"]["avg_sentence_length_words"] == 18.3, "finished work keeps the starting profile"
    short_pieces = measure(["Short one. And another short one. A third."] * 200, {})
    blended = register_numbers(long_notes, short_pieces)["finalized"]
    assert blended["avg_sentence_length_words"] < 18.3, "finished pieces blend into the finalized register"


def test_finished_pieces_are_read_from_their_folder(tmp_path):
    (tmp_path / "essay.md").write_text("# Loneliness\n\nCurleys wife is lonely in a world of men.\n\nShort.\n")
    (tmp_path / "report.txt").write_text("Carbohydrates are our energy source for every athlete.\n")
    assert finished_pieces(tmp_path) == ["Curleys wife is lonely in a world of men.",
                                         "Carbohydrates are our energy source for every athlete."]
    assert finished_pieces(tmp_path / "missing") == []


def test_banned_phrases_are_read_from_either_format():
    old = {"style_profile": {"vocabulary_and_lexical_density": {"banned_ai_cliches": ["delve", "not just X, but Y"]}}}
    new = {"style_profile": {"shared": {"banned_ai_cliches": ["thus", "landscape (figurative)"]}}}
    assert list(banned_patterns(old)) == ["delve"] and list(banned_patterns(new)) == ["thus"]


def test_each_habit_keeps_the_templates_fields():
    template = {"habits": [{"habit": "", "rate": "", "status": ""}]}
    good = {"habits": [{"habit": "Starts with However", "rate": "10%", "status": "established"}]}
    assert format_errors(good, template) == []
    bad = {"habits": [{"habit": "Starts with However", "status": "once"}, "a loose line"]}
    errors = format_errors(bad, template)
    assert any("habits[0]: missing keys ['rate']" in e for e in errors)
    assert any("habits[1]: expected object, found string" in e for e in errors)


def test_the_shipped_template_and_profile_format_agree():
    from pathlib import Path

    template = json.loads(Path("style_profile_template.json").read_text())
    registers = template["style_profile"]["registers"]
    assert list(registers) == ["finalized", "notes"] and registers["finalized"] == registers["notes"]
    assert format_errors(template, template) == []


TEMPLATE = {"style_profile": {
    "syntax_and_burstiness": {"avg_sentence_length_words": 18.3, "min_max_sentence_length": [2, 106],
                              "burstiness_pattern": "High burstiness."},
    "vocabulary_and_lexical_density": {"signature_vocabulary": ["clearly"],
                                       "banned_ai_cliches": ["delve", "landscape (figurative)", "not just X, but Y"]},
}}


def test_format_errors_catch_renamed_retyped_and_reordered_keys():
    assert format_errors(TEMPLATE, TEMPLATE) == []
    grown = json.loads(json.dumps(TEMPLATE))
    grown["style_profile"]["vocabulary_and_lexical_density"]["signature_vocabulary"].append("famously (new 2026-W44)")
    grown["style_profile"]["syntax_and_burstiness"]["burstiness_pattern"] += " Notes run longer."
    assert format_errors(grown, TEMPLATE) == []  # growing inside the format is the point

    broken = json.loads(json.dumps(TEMPLATE))
    syntax = broken["style_profile"]["syntax_and_burstiness"]
    syntax["avg_sentence_length_words"] = "18.6"
    syntax["min_max_sentence_length"] = [1, 50, 106]
    broken["style_profile"]["vocabulary_and_lexical_density"]["signature_vocabulary"].append({"word": "x"})
    errors = format_errors(broken, TEMPLATE)
    assert any("avg_sentence_length_words: expected number" in e for e in errors)
    assert any("expected 2 numbers" in e for e in errors)
    assert any("signature_vocabulary: unexpected object" in e for e in errors)

    sections = TEMPLATE["style_profile"]
    renamed = {"style_profile": {"vocabulary_and_lexical_density": sections["vocabulary_and_lexical_density"],
                                 "syntax": sections["syntax_and_burstiness"]}}
    assert "missing keys ['syntax_and_burstiness'], unexpected keys ['syntax']" in format_errors(renamed, TEMPLATE)[0]
    reordered = {"style_profile": dict(reversed(list(TEMPLATE["style_profile"].items())))}
    assert format_errors(reordered, TEMPLATE) == ["style_profile: keys are out of order"]


@pytest.fixture
def setup(tmp_path):
    vault, state = tmp_path / "vault", tmp_path / "state"
    (vault / "02 The Scrolls").mkdir(parents=True)
    (vault / "80 Scribe Tools").mkdir()
    (vault / "80 Scribe Tools" / "Template.md").write_text("A template line the writer did not write.\n")
    (vault / "00 Review Log.md").write_text("A review written by someone else about the vault.\n")
    template = tmp_path / "template.json"
    template.write_text(json.dumps(TEMPLATE))
    args = ["--vault", str(vault), "--state", str(state), "--profile", str(tmp_path / "profile.json"),
            "--template", str(template)]
    return vault, state, tmp_path / "profile.json", args


def corpus(state):
    return [json.loads(line) for line in (state / "corpus.jsonl").read_text().splitlines()]


def test_runs_stack_week_on_week_and_never_forget(setup):
    vault, state, profile, args = setup
    note = vault / "02 The Scrolls" / "Sparta.md"
    note.write_text('---\ncreated: "13-03-2026"\n---\n'
                    "Sparta had two kings and a council of elders.\n\n"
                    "The helots worked the land for the Spartans who trained for war.\n\n"
                    "Lycurgus wrote the laws which every Spartan had to follow.\n")
    assert main([*args, "--week", "2026-W44"]) == 0
    first = corpus(state)
    assert [e["week"] for e in first] == ["2026-W11"] * 3  # the first run dates by the note's created week
    assert profile.exists() and (state / "history" / "profile_2026-W44.json").exists()

    note.write_text('---\ncreated: "13-03-2026"\n---\n'
                    "Sparta had two kings and a council of elders.\n\n"  # unchanged
                    "The helots worked the land for the Spartans who trained for war all year.\n\n"  # edited
                    "Spartan boys left home at 7 to train in the agoge.\n")  # new; Lycurgus deleted
    assert main([*args, "--week", "2026-W45"]) == 0
    assert len(corpus(state)) == 5
    entries = {e["text"][:20]: e for e in corpus(state) if "superseded" not in e}
    assert entries["Lycurgus wrote the l"]["week"] == "2026-W11"  # deleted from the vault, still remembered
    old, new = (e for e in corpus(state) if e["text"].startswith("The helots"))
    assert old["superseded"] == "2026-W45" and new["edit_of"] == old["id"] and new["week"] == "2026-W45"
    assert entries["Spartan boys left ho"]["week"] == "2026-W45"

    report = json.loads((state / "measurements.json").read_text())
    assert report["corpus"]["paragraphs"] == 4 and report["corpus"]["edited_paragraphs"] == 1
    assert report["corpus"]["kept_after_deletion_from_vault"] == 1
    assert report["this_week"]["paragraphs"] == 2
    assert [w["week"] for w in report["weeks"]] == ["starting profile", "2026-W11", "2026-W45"]
    new_writing = (state / "new_this_week.md").read_text()
    assert "(edited) The helots" in new_writing and "Spartan boys" in new_writing and "Lycurgus" not in new_writing
    assert (state / "history" / "profile_2026-W45.json").exists()
    assert "template line" not in (state / "corpus.jsonl").read_text()  # templates and review logs are skipped

    assert main([*args, "--week", "2026-W45"]) == 0  # a second run in the same week adds nothing
    assert len(corpus(state)) == 5


def test_check_reports_format_problems(setup, capsys):
    _, _, profile, args = setup
    profile.write_text(json.dumps(TEMPLATE))
    assert main([*args, "--check"]) == 0
    profile.write_text(json.dumps({"style_profile": {}}))
    assert main([*args, "--check"]) == 1
    assert "missing keys" in capsys.readouterr().err


def test_finished_pieces_keep_only_the_writers_own_prose_from_any_format(tmp_path):
    from docx import Document

    from update_style_profile import finished_sources, own_prose

    essay = Document()
    essay.add_paragraph("Both poems use vivid imagery, and the contrast gives the scene its meaning here.")
    essay.add_paragraph('Gary writes "With two oranges in my jacket and frost cracking beneath my steps" in winter.')
    essay.add_paragraph("Frost cracking\nbeneath my steps\nmy breath before me\nthen gone")  # a poem's lines
    essay.add_paragraph("Fill in your reflection below in full sentences, using evidence.")  # the template's
    essay.save(tmp_path / "finished work.docx")
    template = Document()
    template.add_paragraph("Fill in your reflection below in full sentences, using evidence.")
    (tmp_path / "templates").mkdir()
    template.save(tmp_path / "templates" / "blank.docx")
    (tmp_path / "notes.md").write_text(
        "Both poems use vivid imagery, and the contrast gives the scene its meaning here.\n\n"
        "Carbohydrates are our energy source, like a battery for every athlete.\n")
    (tmp_path / "~$finished work.docx").write_bytes(b"Word's lock file")
    (tmp_path / "weights.json").write_text('{"notes.md": 0.9}')

    assert finished_sources(tmp_path) == [
        ("finished work.docx", 1.0, ["Both poems use vivid imagery, and the contrast gives the scene its meaning here.",
                                     "Gary writes “…” in winter."]),
        ("notes.md", 0.9, ["Carbohydrates are our energy source, like a battery for every athlete."]),
    ], "quotes cut, verse and template text left out, a repeated paragraph counted once, weights read"
    assert own_prose("y = −0.0075(82.5)2 + 1.2375(82.5) = 51.02 and so on.") is None, "an equation isn't prose"


def test_pdf_lines_run_back_into_paragraphs(monkeypatch):
    import sys
    from types import SimpleNamespace

    from update_style_profile import _pdf_paragraphs

    page = SimpleNamespace(extract_text=lambda: "  2 \nIntroduction: \nThis project models a bridge using \n"
                                                "quadratic functions. \n \nVertex: \n(82.5,52) \n")
    monkeypatch.setitem(sys.modules, "pypdf", SimpleNamespace(PdfReader=lambda data: SimpleNamespace(pages=[page])))
    assert _pdf_paragraphs(b"%PDF") == ["2", "This project models a bridge using quadratic functions.", "(82.5,52)"]


def test_the_vault_ladder_counts_gold_and_silver_notes_toward_finished_work(tmp_path):
    from update_style_profile import LADDER, finished_counts, note_ladder

    assert note_ladder('---\ntags: ["Gold"]\nzk: permanent\n---\nText') == "Gold"
    assert note_ladder("---\ntags:\n  - Silver\n---\nText") == "Silver"
    assert note_ladder('---\ntags: ["Gold"]\nzk: fleeting\n---\nText') == "fleeting", "a capture never counts"
    assert note_ladder("No frontmatter") is None
    sentence = "The polis needed one government strong enough to protect the whole state."
    live = [{"note": "gold.md", "text": sentence}, {"note": "bronze.md", "text": sentence},
            {"note": "silver.md", "text": sentence}]
    ladders = {"gold.md": "Gold", "bronze.md": "Bronze", "silver.md": "Silver"}
    counts, sources = finished_counts([("essay.docx", 1.0, [sentence])], live, ladders, {})
    assert counts["sentences"] == pytest.approx(1 + LADDER["Gold"] + LADDER["Silver"])
    assert sources["files"][0]["file"] == "essay.docx" and set(sources["vault_notes"]) == {"Gold", "Silver"}

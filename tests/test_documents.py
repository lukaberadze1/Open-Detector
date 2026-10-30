import io
import zipfile

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from opendetector.documents import DocumentSection, extract_answer_slots, extract_document, rewritable_blocks


def docx_bytes():
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("   ")
    document.add_paragraph("Second paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "a", "b"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
    slide.shapes.title.text = "Title"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "x", "y"
    presentation.slides.add_slide(presentation.slide_layouts[6])  # blank, so skipped
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_docx_paragraphs_and_table_cells_keep_their_locations():
    assert extract_document("essay.DOCX", docx_bytes()) == [
        DocumentSection("Paragraph 1", "First paragraph.", "docx_paragraph", (0,)),
        DocumentSection("Paragraph 3", "Second paragraph.", "docx_paragraph", (2,)),
        DocumentSection("Table 1, row 1, cell 1", "a", "docx_table_cell", (0, 0, 0)),
        DocumentSection("Table 1, row 1, cell 2", "b", "docx_table_cell", (0, 0, 1)),
    ]


def test_pptx_shapes_and_table_cells_are_separate_sections():
    assert extract_document("deck.pptx", pptx_bytes()) == [
        DocumentSection("Slide 1, Title 1", "Title", "pptx_shape", (0, 0)),
        DocumentSection("Slide 1, Table 1, row 1, cell 1", "x", "pptx_table_cell", (0, 1, 0, 0)),
        DocumentSection("Slide 1, Table 1, row 1, cell 2", "y", "pptx_table_cell", (0, 1, 0, 1)),
    ]


def test_a_pptx_table_reads_as_question_and_answer_pairs():
    """python-pptx's row cells take no negative index, which once crashed every deck with a table."""
    (slot,) = extract_answer_slots("deck.pptx", pptx_bytes())
    assert (slot.question, slot.previous_answer, slot.target.locator) == ("x", "y", (0, 1, 0, 1))


@pytest.mark.parametrize("name", ["essay.txt", "no_extension"])
def test_unsupported_types_are_rejected(name):
    with pytest.raises(ValueError, match=r"Upload a \.md"):
        extract_document(name, b"text")


def test_damaged_or_renamed_files_raise_value_error():
    other_zip = io.BytesIO()
    with zipfile.ZipFile(other_zip, "w") as z:
        z.writestr("notes.txt", "hi")
    for name, data in [("a.docx", b"not a zip"), ("a.pptx", other_zip.getvalue()), ("a.pptx", docx_bytes())]:
        with pytest.raises(ValueError, match=r"isn't a valid \.(docx|pptx) file"):
            extract_document(name, data)


def worksheet_docx_bytes():
    document = Document()
    document.add_heading("Keeping bees", level=1)
    document.add_paragraph("Bees need a dry hive and a nearby source of water.")
    document.add_paragraph("What do bees eat?")
    document.add_paragraph("They live on nectar and pollen from flowers.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Why keep bees", "For the honey."
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_rewriting_a_word_document_skips_headings_questions_and_tables():
    assert [block.text for block in rewritable_blocks("w.docx", worksheet_docx_bytes())] == [
        "Bees need a dry hive and a nearby source of water.",
        "They live on nectar and pollen from flowers.",
    ]


def test_rewriting_slides_skips_titles_and_question_answer_box_pairs():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])  # title only
    slide.shapes.title.text = "Keeping bees"
    for left, top, width, text in ((0.5, 1.5, 9, "A hive needs shelter from the wind."),
                                   (0.5, 5.5, 4, "Describe the hive"), (5.5, 5.5, 4, "A wooden box.")):
        slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(1)).text = text
    buffer = io.BytesIO()
    presentation.save(buffer)
    assert [block.text for block in rewritable_blocks("deck.pptx", buffer.getvalue())] == [
        "A hive needs shelter from the wind.",
    ]


def test_rewriting_markdown_keeps_headings_questions_and_tables():
    text = "# Bees\n\nBees make honey.\n\nQuestion: why?\n\n| a | b |\n|---|---|\n| c | d |\n\nThey sting.\n"
    assert [block.text for block in rewritable_blocks("notes.md", text.encode())] == [
        "Bees make honey.", "They sting.",
    ]

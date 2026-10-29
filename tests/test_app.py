"""The Streamlit app run headless: stand-in engines (no model downloads) against the shipped
combiners, and a stand-in text model on the Red team page writing to a temporary folder."""

import io
import json
import re
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

APP = str(Path(__file__).resolve().parent.parent / "app.py")
HUMAN = "My grandmother kept bees behind the shed, and every summer we got stung at least once. "
AI = "Furthermore, it is important to delve into the multifaceted tapestry of this topic. "


class KeywordClassifier:
    def __init__(self, name):
        self.name = name

    def window_logits(self, texts):
        return np.array([8.0 if "delve" in t else -8.0 for t in texts])


class KeywordBinoculars:
    name = "bino"

    def span_scores(self, text, spans):
        return np.array([0.6 if "delve" in text[a:b] else 1.1 for a, b in spans])


@pytest.fixture(autouse=True)
def stand_in_engines(monkeypatch):
    import opendetector.engines as engines
    from ui import models

    monkeypatch.setattr(engines, "build_engine",
                        lambda name, config: KeywordBinoculars() if name == "bino" else KeywordClassifier(name))
    models.release()  # engines another test loaded
    yield
    models.release()


def button(at, label):
    return next(b for b in at.button if b.label == label)


def analyze(text, profile="full"):
    at = AppTest.from_file(APP, default_timeout=60)
    at.session_state["profile"] = profile
    at.run()
    at.text_area(key="keep:detect:text").input(text)
    button(at, "Analyze").click()
    at.run()
    assert not at.exception, at.exception
    return at, at.session_state["detect:result"]


def test_ai_text_is_flagged_with_every_engine(monkeypatch):
    at, result = analyze(AI * 30)
    assert result["report"].verdict == "likely_ai"
    assert set(result["report"].engines) == {"glyph", "desklib", "bino"}
    assert "likely_ai" in result["bands"]
    assert [tab.label for tab in at.tabs] == ["Highlights", "Sentences", "Details"]
    assert len(result["contributions"]) == 6  # every feature's push, for the Details chart


def test_human_text_on_the_slim_build_runs_desklib_alone():
    _, result = analyze(HUMAN * 30, profile="slim")
    assert result["report"].verdict == "likely_human"
    assert set(result["report"].engines) == {"desklib"}


def test_short_text_abstains():
    at, result = analyze(HUMAN * 3)
    assert result["report"].verdict == "insufficient_text"
    assert any("Need at least 150 words" in info.value for info in at.info)


def test_switched_off_engine_is_left_out():
    at = AppTest.from_file(APP, default_timeout=60)
    at.session_state["profile"] = "full"
    at.session_state["keep:detect:engine:bino"] = False
    at.run()
    at.text_area(key="keep:detect:text").input(AI * 30)
    button(at, "Analyze").click()
    at.run()
    result = at.session_state["detect:result"]
    assert set(result["report"].engines) == {"glyph", "desklib"}
    assert result["missing"] == ["bino_doc", "bino_min"]


def test_redteam_off_leaves_out_the_red_team_page(monkeypatch):
    from ui import nav

    built = []
    monkeypatch.setattr(nav, "redteam_page", lambda: built.append("red team"))
    monkeypatch.setenv("OPENDETECTOR_REDTEAM", "off")
    at = AppTest.from_file(APP, default_timeout=60)
    at.run()
    assert not at.exception, at.exception
    assert built == []


def red_team_page():
    from ui import redteam

    redteam.render()


@pytest.fixture
def stub_model(monkeypatch, tmp_path):
    """The Red team page writing to tmp_path, with a stand-in text model that records its prompts, the
    local model it was offered and its other options, and waits for `gate` before answering. Asked for
    JSON, it answers every block of a group except those in `drop`. Notifications are recorded, and
    nothing keeps the machine awake."""
    from ui import redteam

    for name, path in (("ROOT", tmp_path), ("DATA_ROOT", tmp_path), ("OUTPUT_PATH", tmp_path / "local_redteam.jsonl"),
                       ("EXPORT_DIR", tmp_path / "exports"), ("BATCH_DIR", tmp_path / "batches"),
                       ("STYLE_DIR", tmp_path / "styles"), ("VISITOR_DIR", tmp_path / "visitors")):
        monkeypatch.setattr(redteam, name, path)
    monkeypatch.setattr(redteam, "load_env", lambda: None)  # never read the real .env
    monkeypatch.setattr(redteam.generation_settings, "load_table", lambda: None)  # settings by hand, unless a test
    # gives the priorities a table
    monkeypatch.setattr(redteam, "local_fallback_problem", lambda url, model: None)
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    stub = SimpleNamespace(prompts=[], systems=[], keys=[], urls=[], orders=[], local_models=[], options=[],
                           notifications=[], failures=[], drop=set(), gate=threading.Event())
    stub.gate.set()

    def write(base_url, api_key, models, prompt, temperature, system, local_url="", local_model="", *args,
              **options):
        stub.prompts.append(prompt)
        stub.systems.append(system)
        stub.keys.append(api_key)
        stub.urls.append(base_url)
        stub.orders.append(list(models))
        stub.local_models.append(local_model)
        stub.options.append(options)
        assert stub.gate.wait(30)
        if stub.failures:  # "busy", or "spent:<model>" for a model whose quota runs out
            failure = stub.failures.pop(0)
            if failure.startswith("spent:"):
                args[-1].add(failure.removeprefix("spent:"))
            raise redteam.APIError(f"Every configured text model failed: {failure}")
        if options.get("json_schema"):
            ids = re.findall(r"^\[(B\d+)\]", prompt, re.MULTILINE)
            reply = {"blocks": [{"id": block_id, "text": AI * 15} for block_id in ids if block_id not in stub.drop]}
            return json.dumps(reply), "stub-model", []
        return AI * 15, "stub-model", []

    monkeypatch.setattr(redteam, "generate_with_fallback", write)
    monkeypatch.setattr(redteam, "_notify", stub.notifications.append)
    monkeypatch.setattr(redteam, "_keep_awake", lambda: None)
    yield stub
    stub.gate.set()  # never leave a worker waiting


def run_until_batch_ends(at):
    from ui.redteam import NOTICE

    for _ in range(100):  # the batch runs on a worker thread; each run checks whether it has ended
        if NOTICE in at.session_state:
            break
        time.sleep(0.1)
        at.run()
    assert not at.exception, at.exception
    return at.session_state[NOTICE]


def test_red_team_sample_is_recorded_as_ai_with_its_score(stub_model, tmp_path):
    at = AppTest.from_function(red_team_page, default_timeout=60)
    at.session_state["profile"] = "slim"
    at.run()
    assert not at.exception, at.exception
    at.text_area(key="keep:redteam:prompt").input("An essay on keeping bees.")
    button(at, "Generate sample").click()
    at.run()
    assert not at.exception, at.exception

    rows = [json.loads(line) for line in (tmp_path / "local_redteam.jsonl").read_text().splitlines()]
    assert len(rows) == 1
    assert rows[0] | {"text": None, "id": None, "prompt": None, "group_key": None, "ai_probability": None} == {
        "id": None, "text": None, "label": 1, "dataset": "local_redteam", "domain": "custom_prompt",
        "model": "stub-model", "attack": "prompt+plain", "group": "ai", "group_key": None, "prompt": None,
        "source_file": None, "target_section": None, "ai_probability": None, "verdict": "likely_ai",
        "calibrated": True, "detector_profile": "slim", "style_match": None,
    }
    assert rows[0]["ai_probability"] > 0.9
    assert len(list((tmp_path / "exports").glob("AI_REDTEAM_prompt_*.txt"))) == 1
    assert at.session_state["last_generated_result"]["detect_text"] == rows[0]["text"]


def robustness_page(monkeypatch, replies, until_target):
    """The Red team page in robustness mode, its text model answering `replies` in turn (text, model) and
    its draft loop replaced by `until_target`, which these tests own."""
    from ui import redteam

    answers = iter(replies)
    monkeypatch.setattr(redteam, "generate_with_fallback", lambda *args, **options: (*next(answers), []))
    monkeypatch.setattr(redteam, "generate_until_target", until_target)
    at = AppTest.from_function(red_team_page, default_timeout=60)
    at.session_state["profile"] = "slim"
    at.session_state["keep:redteam:evasion"] = True
    at.run()
    at.text_area(key="keep:redteam:prompt").input("An essay on keeping bees.")
    button(at, "Generate sample").click()
    at.run()
    assert not at.exception, at.exception
    return at


def test_robustness_mode_keeps_the_lowest_draft_and_saves_every_one(stub_model, tmp_path, monkeypatch):
    prompts = []

    def lowest_of_three(generate, score, prompt):
        prompts.append(prompt)
        drafts = [generate(prompt) for _ in range(3)]
        return SimpleNamespace(text=min(drafts, key=lambda draft: score(draft)["ai_probability"]))

    replies = [(AI * 15, "model-a"), (HUMAN * 15, "model-b"), (AI * 20, "model-c")]
    at = robustness_page(monkeypatch, replies, lowest_of_three)
    assert prompts[0].startswith("Write an AI-generated red-team sample")

    rows = [json.loads(line) for line in (tmp_path / "local_redteam.jsonl").read_text().splitlines()]
    assert [(r["model"], r["evasion_kept"], r["label"]) for r in rows] == [
        ("model-b", True, 1), ("model-a", False, 1), ("model-c", False, 1),
    ], "the kept draft first, credited to the model that wrote it, then the others"
    assert rows[0]["text"] == (HUMAN * 15).strip()
    assert rows[0]["evasion_attempts"] == 3 and rows[0]["evasion_reached_target"]
    assert rows[0]["verdict"] == "likely_human" and rows[1]["verdict"] == rows[2]["verdict"] == "likely_ai"
    assert len({r["group_key"] for r in rows}) == 1 and len({r["id"] for r in rows}) == 3
    assert at.session_state["last_generated_result"]["detect_text"] == rows[0]["text"]


def test_robustness_mode_keeps_its_drafts_when_the_models_run_out(stub_model, tmp_path, monkeypatch):
    from ui import redteam

    def quota_runs_out(generate, score, prompt):
        score(generate(prompt))
        raise redteam.APIError("Every configured text model failed: quota")

    at = robustness_page(monkeypatch, [(AI * 15, "model-a")], quota_runs_out)
    assert any("stopped after 1 draft" in w.value for w in at.warning)
    rows = [json.loads(line) for line in (tmp_path / "local_redteam.jsonl").read_text().splitlines()]
    assert [(r["model"], r["evasion_kept"], r["evasion_reached_target"]) for r in rows] == [("model-a", True, False)]


def test_robustness_mode_turns_scoring_on(stub_model, tmp_path, monkeypatch):
    from ui import redteam

    scores = []

    def one_draft(generate, score, prompt):
        text = generate(prompt)
        scores.append(score(text))
        return SimpleNamespace(text=text)

    monkeypatch.setattr(redteam, "generate_until_target", one_draft)
    at = AppTest.from_function(red_team_page, default_timeout=60)
    at.session_state["profile"] = "slim"
    at.session_state["keep:redteam:evasion"] = True
    at.session_state["keep:redteam:score"] = False  # robustness mode needs scores, so it scores anyway
    at.run()
    at.text_area(key="keep:redteam:prompt").input("An essay on keeping bees.")
    button(at, "Generate sample").click()
    at.run()
    assert not at.exception, at.exception
    assert len(scores) == 1 and scores[0] is not None, "each draft was scored"
    row = json.loads((tmp_path / "local_redteam.jsonl").read_text())
    assert row["evasion_kept"] and row["verdict"] == "likely_ai"


def test_red_team_writes_in_the_style_profile_picked(stub_model, tmp_path):
    styles = tmp_path / "styles"
    styles.mkdir()
    (styles / "alexandria_profile.json").write_text(json.dumps({"tone": "dry"}))
    (styles / "friend.json").write_text(json.dumps({"style_profile": {"tone": "breezy"}}))
    at = AppTest.from_function(red_team_page, default_timeout=60)
    at.session_state["profile"] = "slim"
    at.run()
    picker = at.selectbox(key="keep:redteam:style_name")
    assert (picker.options, picker.value) == (["alexandria_profile", "friend"], "alexandria_profile")
    picker.select("friend")
    at.text_area(key="keep:redteam:prompt").input("An essay on keeping bees.")
    button(at, "Generate sample").click()
    at.run()
    assert not at.exception, at.exception

    assert "tone: breezy" in stub_model.systems[-1]
    row = json.loads((tmp_path / "local_redteam.jsonl").read_text())
    assert row["attack"] == "prompt+friend"


def essay_docx():
    document = Document()
    document.add_heading("Keeping bees", level=1)
    document.add_paragraph("Bees need a dry hive and a nearby source of water.")
    document.add_paragraph("What do bees eat?")
    document.add_paragraph("They live on nectar and pollen from flowers.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def word_document_page():
    from ui.components import Upload

    at = AppTest.from_function(red_team_page, default_timeout=60)
    at.session_state["profile"] = "slim"
    at.session_state["upload:redteam:file"] = Upload("bees.docx", essay_docx())
    at.session_state["keep:redteam:task"] = "blocks"
    at.run()
    return at


def test_red_team_rewrites_every_block_of_a_word_document(stub_model, tmp_path):
    at = word_document_page()
    assert not at.exception, at.exception
    assert any("Fill answer boxes** isn't available for this file: found no question/answer pairs" in caption.value
               for caption in at.caption)

    button(at, "Rewrite 2 blocks").click()
    at.run()
    assert run_until_batch_ends(at) == ("success", "All 2 blocks of bees.docx are done and exported.")
    assert [toast.value for toast in at.toast] == ["All 2 blocks of bees.docx are done and exported."]
    assert stub_model.notifications == ["All 2 blocks of bees.docx are written and exported."]

    (prompt,) = stub_model.prompts  # both blocks in one request, which sends the document once
    assert prompt.count("Full document context:") == 1
    assert "[B1] A rewrite of the existing text block labelled Paragraph 2. Aim for about 11 words." in prompt
    assert "[B2] A rewrite of the existing text block labelled Paragraph 4" in prompt
    assert stub_model.options[0]["thinking"] == "low" and stub_model.options[0]["json_schema"]
    (exported,) = (tmp_path / "exports").glob("AI_REDTEAM_full_document_*.docx")
    assert [p.text for p in Document(exported).paragraphs] == [
        "Keeping bees", (AI * 15).strip(), "What do bees eat?", (AI * 15).strip(),
    ]
    rows = [json.loads(line) for line in (tmp_path / "local_redteam.jsonl").read_text().splitlines()]
    assert [(row["domain"], row["attack"], row["label"]) for row in rows] == [
        ("document_text_block", "full_document+plain", 1),
        ("document_text_block", "full_document+plain", 1),
        ("document_combined", "full_document+plain", 1),
    ]
    assert all(row["verdict"] == "likely_ai" for row in rows)
    assert not list((tmp_path / "batches").glob("*.json")), "a finished batch's checkpoint is discarded"


def test_a_running_batch_takes_over_the_page_with_a_countdown(stub_model, tmp_path):
    stub_model.gate.clear()  # hold the first block until the end of the test
    at = word_document_page()
    button(at, "Rewrite 2 blocks").click()
    at.run()
    assert not at.exception, at.exception

    def progress_shown():
        return any("Writing bees.docx** · Rewrite every block" in m.value for m in at.markdown)

    assert progress_shown()
    assert {m.label: m.value for m in at.metric} == {
        "Time left": "Timing the first block…", "Done around": "Soon", "Blocks written": "0 of 2",
    }
    assert [b.label for b in at.button] == ["Stop"], "while a batch runs, nothing else can start"

    # Whatever the page is set to, it shows only the running batch.
    at.session_state["keep:redteam:match"] = False
    at.session_state["keep:redteam:task"] = "sample"
    at.run()
    assert progress_shown()
    assert [b.label for b in at.button] == ["Stop"]

    # It ends while the page is set to something else, and is exported with its own settings.
    stub_model.gate.set()
    assert run_until_batch_ends(at)[0] == "success"
    assert len(list((tmp_path / "exports").glob("AI_REDTEAM_full_document_*.docx"))) == 1
    assert "Aim for about 11 words" in stub_model.prompts[0], "the batch kept its match-the-length setting"


def public_page(visitor):
    """The Red team page on a public deploy, as the browser session `visitor` sees it."""
    at = AppTest.from_function(red_team_page, default_timeout=60)
    at.session_state["profile"] = "slim"
    at.session_state["redteam:visitor"] = visitor
    at.run()
    assert not at.exception, at.exception
    return at


def test_public_mode_is_on_by_default_on_streamlit_community_cloud(monkeypatch):
    from ui import redteam

    monkeypatch.delenv("OPENDETECTOR_PUBLIC", raising=False)
    monkeypatch.setattr(redteam, "ROOT", Path("/mount/src/open-detector"))
    assert redteam.public()
    monkeypatch.setenv("OPENDETECTOR_PUBLIC", "off")
    assert not redteam.public()
    monkeypatch.delenv("OPENDETECTOR_PUBLIC")
    monkeypatch.setattr(redteam, "ROOT", Path("/Users/someone/Open-Detector"))
    assert not redteam.public()


def test_public_mode_writes_only_with_the_visitors_own_key(stub_model, monkeypatch):
    from eval.generate_essays import GEMINI_API_URL

    monkeypatch.setenv("OPENDETECTOR_PUBLIC", "on")  # the host's GEMINI_API_KEY is set, and must go unused
    at = public_page("a")
    assert not any(t.key == "keep:redteam:base_url" for t in at.text_input), "no visitor-chosen API address"
    assert not any(t.key == "keep:redteam:evasion" for t in at.toggle), "no detector-guided drafting for strangers"
    at.text_area(key="keep:redteam:prompt").input("An essay on keeping bees.")
    button(at, "Generate sample").click()
    at.run()
    assert [e.value for e in at.error] == ["Paste your Gemini API key in the sidebar first."]
    assert stub_model.prompts == []

    at.text_input(key="keep:redteam:api_key").input("visitor-key")
    button(at, "Generate sample").click()
    at.run()
    assert not at.exception, at.exception
    assert (stub_model.keys, stub_model.urls, stub_model.local_models) == (["visitor-key"], [GEMINI_API_URL], [""])


def test_public_visitors_see_only_their_own_profiles_and_samples(stub_model, monkeypatch, tmp_path):
    monkeypatch.setenv("OPENDETECTOR_PUBLIC", "on")
    (tmp_path / "styles").mkdir()
    (tmp_path / "styles" / "alexandria_profile.json").write_text("{}")  # the host's own, never offered
    own_styles = tmp_path / "visitors" / "a" / "styles"
    own_styles.mkdir(parents=True)
    (own_styles / "friend.json").write_text(json.dumps({"tone": "breezy"}))
    at = public_page("a")
    assert at.selectbox(key="keep:redteam:style_name").options == ["friend"]
    at.text_input(key="keep:redteam:api_key").input("visitor-key")
    at.text_area(key="keep:redteam:prompt").input("An essay on keeping bees.")
    button(at, "Generate sample").click()
    at.run()
    assert not at.exception, at.exception
    assert "tone: breezy" in stub_model.systems[-1]
    assert len((tmp_path / "visitors" / "a" / "local_redteam.jsonl").read_text().splitlines()) == 1
    assert not (tmp_path / "local_redteam.jsonl").exists()

    other = public_page("b")
    assert "last_generated_result" not in other.session_state
    assert not any(s.key == "keep:redteam:style_name" for s in other.selectbox)
    assert any("No samples yet" in i.value for i in other.info)


def test_public_visitors_dont_see_each_others_batches(stub_model, monkeypatch):
    from ui.components import Upload

    monkeypatch.setenv("OPENDETECTOR_PUBLIC", "on")
    stub_model.gate.clear()  # hold the first block until the end of the test
    at = AppTest.from_function(red_team_page, default_timeout=60)
    at.session_state["profile"] = "slim"
    at.session_state["redteam:visitor"] = "a"
    at.session_state["keep:redteam:api_key"] = "visitor-key"
    at.session_state["upload:redteam:file"] = Upload("bees.docx", essay_docx())
    at.session_state["keep:redteam:task"] = "blocks"
    at.run()
    button(at, "Rewrite 2 blocks").click()
    at.run()
    assert [b.label for b in at.button] == ["Stop"]

    other = public_page("b")
    assert "Generate sample" in [b.label for b in other.button], "another visitor's batch doesn't take over"
    stub_model.gate.set()
    assert run_until_batch_ends(at)[0] == "success"


def test_an_added_style_profile_is_kept_for_the_visitor_and_picked(monkeypatch, tmp_path):
    from ui import redteam

    state = {"redteam:style_upload": SimpleNamespace(name="My Style.json", getvalue=lambda: b'{"tone": "breezy"}')}
    monkeypatch.setattr(redteam, "st", SimpleNamespace(session_state=state))
    monkeypatch.setattr(redteam, "VISITOR_DIR", tmp_path / "visitors")
    monkeypatch.setenv("OPENDETECTOR_PUBLIC", "on")
    redteam._add_style_profile()
    saved = tmp_path / "visitors" / state["redteam:visitor"] / "styles" / "My_Style.json"
    assert saved.read_text() == '{"tone": "breezy"}'
    assert (state["keep:redteam:style_name"], state["redteam:style_notice"][0]) == ("My_Style", "success")

    for name, data in (("list.json", b"[1, 2]"), ("broken.json", b"{"), ("My Style.json", b'{"tone": "dry"}')):
        state["redteam:style_upload"] = SimpleNamespace(name=name, getvalue=lambda data=data: data)
        redteam._add_style_profile()
        assert state["redteam:style_notice"][0] == "error", name
    assert saved.read_text() == '{"tone": "breezy"}', "a different profile of the same name isn't overwritten"


def test_block_prompts_start_with_the_same_document():
    """A local model keeps a prompt's start cached, so every block's prompt opens with the document."""
    from eval.redteam_batch import block_prompt
    from opendetector.documents import extract_answer_slots, extract_document, rewritable_blocks
    from ui import redteam
    from ui.components import Upload

    data = essay_docx()
    plan = redteam.Plan("Keep it short.", Upload("bees.docx", data), extract_document("bees.docx", data),
                        extract_answer_slots("bees.docx", data), rewritable_blocks("bees.docx", data), "blocks",
                        "slim", rewrite_all_blocks=True)
    prompts = [block_prompt(plan, job) for job in redteam._batch_jobs(plan)]
    shared = f"Full document context:\n{plan.context_text}\n\n---\n\nRewrite only the existing text block labelled "
    assert [prompt.startswith(shared) for prompt in prompts] == [True, True]
    assert [prompt.endswith(previous) for prompt, previous in zip(prompts, (
        "Bees need a dry hive and a nearby source of water.", "They live on nectar and pollen from flowers."),
        strict=True)] == [True, True]


def test_a_local_model_too_big_for_the_machine_is_left_out(stub_model, monkeypatch):
    from ui import models, redteam

    monkeypatch.setattr(models, "RAM_GB", 8.0)
    monkeypatch.setattr(redteam, "SMALL_MACHINE_LOCAL_CONTEXT", 4096)  # smaller than any prompt's window
    at = word_document_page()
    button(at, "Rewrite 2 blocks").click()
    at.run()
    run_until_batch_ends(at)
    assert stub_model.local_models == [""], "gemma3:4b, the page's default, is never offered"


def test_a_batch_waits_out_busy_models_and_drops_only_spent_ones(stub_model, monkeypatch):
    from ui import redteam

    monkeypatch.setattr(redteam, "BUSY_WAIT", 0)
    stub_model.failures = ["spent:gemini-3.8-flash", "busy"]
    at = word_document_page()
    at.session_state["keep:redteam:group_size"] = 1
    at.session_state["keep:redteam:parallel"] = 1
    button(at, "Rewrite 2 blocks").click()
    at.run()
    assert run_until_batch_ends(at)[0] == "success"
    assert len(stub_model.prompts) == 4, "the first block is asked three times, the second once"
    first, second, third, _ = stub_model.orders
    assert "gemini-3.8-flash" in first and "gemini-3.8-flash" not in second
    assert third == second, "a model that was only busy stays in the list"


def test_a_block_a_grouped_reply_leaves_out_is_written_on_its_own(stub_model, tmp_path):
    stub_model.drop = {"B2"}
    at = word_document_page()
    button(at, "Rewrite 2 blocks").click()
    at.run()
    assert run_until_batch_ends(at)[0] == "success"
    assert len(stub_model.prompts) == 2
    assert stub_model.prompts[1].split("---\n\n")[1].startswith(
        "Rewrite only the existing text block labelled Paragraph 4")
    assert "json_schema" not in stub_model.options[1]
    (exported,) = (tmp_path / "exports").glob("AI_REDTEAM_full_document_*.docx")
    assert [p.text for p in Document(exported).paragraphs][1::2] == [(AI * 15).strip()] * 2


def test_the_thinking_level_picked_reaches_the_model(stub_model):
    at = AppTest.from_function(red_team_page, default_timeout=60)
    at.session_state["profile"] = "slim"
    at.run()
    at.select_slider(key="keep:redteam:thinking").set_value("minimal")
    at.text_area(key="keep:redteam:prompt").input("An essay on keeping bees.")
    button(at, "Generate sample").click()
    at.run()
    assert not at.exception, at.exception
    assert stub_model.options == [{"thinking": "minimal"}]


def test_robustness_mode_batches_draft_one_block_per_request(stub_model):
    at = word_document_page()
    at.toggle(key="keep:redteam:evasion").set_value(True)
    at.run()
    button(at, "Rewrite 2 blocks").click()
    at.run()
    assert run_until_batch_ends(at)[0] == "success"
    assert stub_model.prompts and not any(options.get("json_schema") for options in stub_model.options)
    assert {prompt.split("---\n\n")[1].split(". ")[0] for prompt in stub_model.prompts} == {
        "Rewrite only the existing text block labelled Paragraph 2",
        "Rewrite only the existing text block labelled Paragraph 4",
    }


def measured_table(monkeypatch):
    """The priorities reading a measured table: a careful setting, a balanced one, and a quick one."""
    from ui import redteam

    def row(model, thinking, group, ai, style, seconds, requests, tokens):
        return {"model": model, "thinking": thinking, "group_size": group, "ai": ai, "style": style,
                "seconds_per_block": seconds, "requests_per_block": requests, "tokens_per_block": tokens,
                "measured": True}

    table = {"context_chars": 1000, "daily_requests": {"gemini-3.6-flash": 20}, "rows": [
        row("gemini-3.8-flash", "high", 1, 0.20, 80, 30.0, 1.0, 9000),
        row("gemini-3.6-flash", "low", 8, 0.35, 70, 3.0, 0.125, 1500),
        row("gemini-3.5-flash-lite", "minimal", 8, 0.80, 30, 1.0, 0.125, 1200),
    ]}
    monkeypatch.setattr(redteam.generation_settings, "load_table", lambda: table)


def test_points_pick_the_batch_settings(stub_model, monkeypatch):
    measured_table(monkeypatch)
    at = word_document_page()
    assert not any(s.key == "keep:redteam:group_size" for s in at.slider), "points, not sliders"
    assert any("gemini-3.6-flash" in m.value and "8 blocks per request · 2 at once" in m.value for m in at.markdown)
    assert any("1 request (5% of its 20 free requests a day)" in c.value for c in at.caption)
    for measure in ("detector", "style", "quota"):  # every point on speed
        for _ in range(3):
            at.button(key=f"redteam:less:{measure}").click().run()
    for _ in range(3):
        at.button(key="redteam:more:speed").click().run()
    assert at.button(key="redteam:more:speed").disabled, "six points is the most one priority takes"
    assert any("gemini-3.5-flash-lite" in m.value and "4 at once" in m.value for m in at.markdown)
    button(at, "Rewrite 2 blocks").click()
    at.run()
    assert run_until_batch_ends(at)[0] == "success"
    assert stub_model.options[0]["thinking"] == "minimal"
    assert stub_model.orders[0][0] == "gemini-3.5-flash-lite", "the batch starts with the model the points picked"


def test_settings_can_still_be_set_by_hand(stub_model, monkeypatch):
    measured_table(monkeypatch)
    at = word_document_page()
    at.toggle(key="keep:redteam:by_hand").set_value(True).run()
    assert at.slider(key="keep:redteam:group_size").value == 6
    at.select_slider(key="keep:redteam:thinking").set_value("high").run()
    button(at, "Rewrite 2 blocks").click()
    at.run()
    assert run_until_batch_ends(at)[0] == "success"
    assert stub_model.options[0]["thinking"] == "high" and stub_model.orders[0][0] == "gemini-3.8-flash"


def test_the_downloadable_app_offers_robustness_mode_but_no_local_model(stub_model, monkeypatch):
    monkeypatch.setenv("OPENDETECTOR_PACKAGED", "on")
    at = AppTest.from_function(red_team_page, default_timeout=60)
    at.session_state["profile"] = "slim"
    at.run()
    assert not at.exception, at.exception
    assert any(t.key == "keep:redteam:evasion" for t in at.toggle)
    assert at.text_input(key="keep:redteam:local_model").value == ""


def test_the_guide_opens_first_once_then_detect_does(monkeypatch, tmp_path):
    from ui import prefs

    monkeypatch.setattr(prefs, "PATH", tmp_path / "ui_prefs.json")  # a computer the app hasn't run on
    at = AppTest.from_file(APP, default_timeout=60).run()
    assert not at.exception, at.exception
    assert "1 · Check a text" in [s.value for s in at.subheader]
    assert prefs.load()["welcomed"]

    at = AppTest.from_file(APP, default_timeout=60).run()
    assert not at.exception, at.exception
    assert "1 · Check a text" not in [s.value for s in at.subheader]
    assert any(t.key == "keep:detect:text" for t in at.text_area)


def test_public_mode_has_no_guide(monkeypatch, tmp_path):
    from ui import nav, prefs

    monkeypatch.setattr(prefs, "PATH", tmp_path / "ui_prefs.json")
    monkeypatch.setenv("OPENDETECTOR_PUBLIC", "on")
    monkeypatch.setattr(nav, "guide_page", lambda: pytest.fail("a public deploy built the Guide page"))
    at = AppTest.from_file(APP, default_timeout=60).run()
    assert not at.exception, at.exception
    assert not prefs.PATH.exists(), "a public deploy remembers nothing"


def guide_page():
    from ui import guide

    guide.render()


@pytest.fixture
def key_file(monkeypatch, tmp_path):
    """The Guide page saving keys to tmp_path/.env, with Google accepting only "good-key"."""
    from functools import partial

    from opendetector import env
    from ui import guide

    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.setattr(guide, "load_env", lambda: None)  # never read the real .env
    monkeypatch.setattr(guide, "save_env", partial(env.save_env, path=tmp_path / ".env"))
    monkeypatch.setattr(guide, "check_key", lambda key: (True, None) if key == "good-key" else
                        (False, "Google didn't accept this key."))
    return tmp_path / ".env"


def save_key(at, key):
    at.text_input[0].input(key)
    button(at, "Save key" if "Save key" in [b.label for b in at.button] else "Replace key").click()
    return at.run()


def test_the_guide_saves_a_key_google_accepts_and_no_other(key_file):
    at = AppTest.from_function(guide_page, default_timeout=60).run()
    assert not at.exception, at.exception
    save_key(at, "wrong-key")
    assert [e.value for e in at.error] == ["Google didn't accept this key."]
    assert not key_file.exists()

    save_key(at, " good-key ")
    assert not at.exception, at.exception
    assert key_file.read_text() == "GEMINI_API_KEY=good-key\n"
    assert any("ending in …-key" in s.value for s in at.success)

    button(at, "Remove the key").click()
    at.run()
    assert key_file.read_text() == ""

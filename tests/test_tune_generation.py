"""The generation-settings benchmark, and the priorities that pick a setting from its table, with a
stand-in Gemini and detector."""

import io
import json
import random
import re
from types import SimpleNamespace

import pytest
from docx import Document

from eval import generate_essays, generation_settings, tune_generation
from eval.generation_settings import EVEN, MEASURES


def row(model, thinking, group, ai, style, seconds, requests, tokens):
    return {"model": model, "thinking": thinking, "group_size": group, "ai": ai, "ai_sd": None, "style": style,
            "style_sd": None, "seconds_per_block": seconds, "requests_per_block": requests,
            "tokens_per_block": tokens, "trials": 1, "measured": True}


TABLE = {"context_chars": 10_000, "rows": [
    row("careful", "high", 1, 0.20, 80, 30.0, 1.0, 9000),  # best past the detector and in style; slow and costly
    row("careful", "low", 8, 0.35, 70, 3.0, 0.125, 1500),  # a little of each
    row("quick", "minimal", 8, 0.80, 30, 1.0, 0.125, 1200),  # fastest and cheapest, and easily caught
]}


def alone(measure, **options):
    return generation_settings.choose(TABLE, {m: 6 if m == measure else 0 for m in MEASURES}, **options)


def test_points_spread_evenly_pick_the_setting_with_a_little_of_each():
    setting = generation_settings.choose(TABLE, EVEN)
    assert (setting.model, setting.thinking, setting.group_size, setting.parallel) == ("careful", "low", 8, 2)


def test_all_points_on_one_measure_pick_its_best_setting():
    assert alone("detector").thinking == alone("style").thinking == "high"
    assert alone("speed").model == alone("quota").model == "quick"
    assert alone("speed").parallel == 4 and alone("quota").parallel == 1, "only speed points add requests at once"


def test_the_pick_keeps_to_the_models_on_offer_and_to_one_block_for_a_single_sample():
    assert alone("speed", models=["careful"]).thinking == "low"
    single = alone("speed", batch=False)
    assert (single.group_size, single.parallel) == (1, 1)
    assert generation_settings.choose(TABLE, EVEN, models=["other"]) is None


def test_an_unmeasured_style_counts_the_same_for_every_setting():
    table = {**TABLE, "rows": [{**r, "style": None} for r in TABLE["rows"]]}
    assert generation_settings.choose(table, {"style": 6}).model == "careful", "ties go to fewer requests"
    assert generation_settings.choose(table, {"style": 6}).group_size == 8


def test_a_gap_of_a_few_points_counts_as_a_few_points():
    rows = [row("a", "high", 8, 0.02, 60, 9.0, 0.2, 3000), row("a", "minimal", 8, 0.07, 60, 3.0, 0.2, 3000)]
    assert generation_settings.choose({"rows": rows}, EVEN).thinking == "minimal", (
        "5 points of AI likelihood weigh 5 points, not the whole range, against three times the speed")


def test_a_model_with_a_small_daily_allowance_loses_out_on_a_large_batch():
    rows = [row("capped", "minimal", 8, 0.03, 45, 0.8, 0.125, 2000),  # a little better, 20 requests a day
            row("open", "minimal", 8, 0.05, 40, 0.9, 0.125, 2000)]
    table = {"rows": rows, "daily_requests": {"capped": 20}}
    assert generation_settings.choose(table, EVEN, blocks=10).model == "capped", "10 blocks take 1 of its 20"
    assert generation_settings.choose(table, EVEN, blocks=160).model == "open", "160 blocks would take all 20"
    assert generation_settings.choose({"rows": rows}, EVEN, blocks=160).model == "capped", "no limit known, no cost"
    estimate = generation_settings.estimate(generation_settings.choose(table, EVEN, blocks=10), 10, table)
    assert estimate["daily_limit"] == 20


def test_estimates_scale_with_the_document():
    setting = generation_settings.choose(TABLE, EVEN)
    estimate = generation_settings.estimate(setting, 99, TABLE, context_chars=20_000)
    assert estimate == {"seconds": 99 * 3.0 / 2, "requests": 99 * 0.125, "tokens": 99 * 1500 * 2.0,
                        "daily_limit": None}


def test_only_the_input_a_request_carries_scales_with_the_document():
    table = {"context_chars": 50_000, "rows": [{**row("m", "minimal", 1, 0.1, 50, 5.0, 1.0, 16_000),
                                                "input_tokens_per_block": 15_000}]}
    setting = generation_settings.choose(table, EVEN, batch=False)
    bare = generation_settings.estimate(setting, 1, table, context_chars=0)
    assert bare["tokens"] == pytest.approx(15_000 * 2000 / 50_000 + 1000), (
        "a sample with no profile or document still sends its instructions, and writes as much as before")
    assert generation_settings.estimate(setting, 1, table, context_chars=100_000)["tokens"] == pytest.approx(31_000)


def test_a_missing_or_empty_table_is_none(tmp_path):
    assert generation_settings.load_table(tmp_path / "missing.json") is None
    (tmp_path / "empty.json").write_text('{"rows": []}')
    assert generation_settings.load_table(tmp_path / "empty.json") is None


def test_the_main_design_measures_each_model_once_per_level_and_the_groups_on_one():
    trials = tune_generation.planned_trials(["gemini-3.8-flash", "b"], ["minimal", "low"], [1, 4], 2, "main", "b")
    settings = sorted({(t["model"], t["thinking"], t["group_size"]) for t in trials})
    assert settings == [("b", "low", 1), ("b", "low", 4), ("b", "minimal", 4), ("gemini-3.8-flash", "low", 4)], (
        "3.8 Flash has no minimal level, so it is asked for low once")
    assert len(trials) == 2 * len(settings)
    full = tune_generation.planned_trials(["a"], ["low", "high"], [1, 4], 1, "full", "a")
    assert len(full) == 4


def document(blocks):
    doc = Document()
    doc.add_heading("Field notes", level=1)
    for text in blocks:
        doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


BLOCKS = [f"Paragraph {n} says the hive needs water and shade in the summer months." for n in range(1, 7)]


class FakeGemini:
    """Stands in for generate_essays.chat: a group's JSON reply leaves out the blocks in `drop`, the
    first `busy` requests fail with a 503, and every answer reports 1,000 tokens."""

    def __init__(self, monkeypatch, drop=(), busy=0):
        self.calls, self.drop, self.busy = [], set(drop), busy
        monkeypatch.setattr(generate_essays, "chat", self.chat)
        monkeypatch.setattr(tune_generation.time, "sleep", lambda seconds: None)

    def chat(self, base_url, api_key, model, prompt, temperature, retries=5, system=None, timeout=300,
             thinking=None, json_schema=None, usage=None):
        self.calls.append({"model": model, "thinking": thinking, "json": json_schema is not None, "system": system})
        if self.busy:
            self.busy -= 1
            raise generate_essays.APIError("HTTP 503 from Gemini Interactions API: high demand")
        usage.update({"total_tokens": 1000, "total_input_tokens": 900, "total_output_tokens": 100})
        if json_schema:
            ids = [i for i in re.findall(r"^\[(B\d+)\]", prompt, re.MULTILINE) if i not in self.drop]
            return json.dumps({"blocks": [{"id": i, "text": f"New text {i}."} for i in ids]})
        return "A block written on its own."


def test_a_trial_rewrites_the_run_as_a_batch_would(monkeypatch):
    gemini = FakeGemini(monkeypatch, drop={"B2"}, busy=1)
    jobs, plan = tune_generation.sample_jobs("notes.docx", document(BLOCKS), words=40)
    assert len(jobs) == 4, "blocks from the first until the run holds 40 words"
    trial = {"model": "m", "thinking": "low", "group_size": 3, "repeat": 1, "key": "m|low|3|1"}
    result = tune_generation.run_trial(trial, jobs, plan, "Style rules.", "key", 0.8)
    assert [call["json"] for call in gemini.calls] == [True, True, False, False], (
        "a group that answered busy is asked again; its left-out block goes alone, and so does the last block")
    assert (result["requests"], result["retried"], result["left_out"]) == (3, 1, 1)
    assert result["tokens"]["total_tokens"] == 3000 and result["blocks"] == 4
    assert result["text"].split("\n\n") == ["New text B1.", "A block written on its own.", "New text B3.",
                                            "A block written on its own."]
    assert all(call["system"] == "Style rules." and call["thinking"] == "low" for call in gemini.calls)


def test_settings_not_measured_are_predicted_from_the_group_size_effect():
    def trial(model, thinking, group, ai, style, seconds, requests):
        return {"model": model, "thinking": thinking, "group_size": group, "blocks": 10, "ai": ai, "style": style,
                "seconds": seconds, "requests": requests, "tokens": {"total_tokens": 100 * requests}}

    rows = tune_generation.fit([
        trial("base", "low", 8, 0.40, 60.0, 20.0, 2),
        trial("base", "low", 1, 0.30, 70.0, 50.0, 10),
        trial("other", "high", 8, 0.20, 50.0, 40.0, 2),
    ], baseline="base")
    predicted = next(r for r in rows if (r["model"], r["group_size"]) == ("other", 1))
    assert not predicted["measured"]
    assert predicted["ai"] == pytest.approx(0.10) and predicted["style"] == pytest.approx(60.0), "effects add"
    assert predicted["seconds_per_block"] == pytest.approx(10.0)
    assert predicted["requests_per_block"] == pytest.approx(1.0)
    assert len(rows) == 4


WORDS = ["the", "hive", "bees", "water", "shade", "honey", "frame", "comb", "queen", "smoke", "field", "flower",
         "summer", "garden", "and", "a", "of", "in"]


def plain_paragraph(rng):
    """Three to five short plain sentences: the writer the tests' corpus is made of."""
    return " ".join(" ".join(rng.choice(WORDS) for _ in range(rng.randint(5, 12))).capitalize() + "."
                    for _ in range(rng.randint(3, 5)))


def corpus_profile(tmp_path, paragraphs=60):
    """A style profile with a corpus of short plain sentences beside it, like alexandria's."""
    styles = tmp_path / "styles"
    (styles / "writer").mkdir(parents=True)
    (styles / "writer_profile.json").write_text("{}")
    rng = random.Random(0)
    lines = [json.dumps({"text": plain_paragraph(rng), "week": "2026-W40", "id": str(n)}) for n in range(paragraphs)]
    (styles / "writer" / "corpus.jsonl").write_text("\n".join(lines) + "\n")
    return styles / "writer_profile.json"


def test_style_match_ranks_a_text_among_the_writers_own_passages(tmp_path):
    yardstick = tune_generation.StyleYardstick.build(corpus_profile(tmp_path), words=120)
    assert yardstick is not None and len(yardstick.passages) >= 5
    rng = random.Random(1)
    alike = [yardstick.match("\n\n".join(plain_paragraph(rng) for _ in range(4))) for _ in range(10)]
    assert 30 < sum(alike) / len(alike) < 70, "passages like the writer's own rank near the middle, on average"
    unlike = ("Notwithstanding the multifaceted considerations — economic, ecological; aesthetic — which "
              "invariably accompany apicultural undertakings, practitioners must, if they are to flourish, "
              "cultivate meticulous attentiveness to environmental variability; ") * 4
    assert yardstick.match(unlike) == 0, "further from the profile than any of the writer's passages"
    assert tune_generation.StyleYardstick.build(tmp_path / "styles" / "plain.json", words=120) is None, (
        "no corpus beside the profile, no style match")


def test_a_run_measures_scores_writes_the_table_and_resumes(monkeypatch, tmp_path):
    gemini = FakeGemini(monkeypatch)
    monkeypatch.setattr(tune_generation, "load_env", lambda: None)
    monkeypatch.setattr(tune_generation, "LABELLED", tmp_path / "tuning_redteam.jsonl")
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    import opendetector

    detector = SimpleNamespace(analyze=lambda text, sentences: SimpleNamespace(ai_probability=0.4))
    monkeypatch.setattr(opendetector.Detector, "from_config", classmethod(lambda cls, config: detector))
    source = tmp_path / "notes.docx"
    source.write_bytes(document(BLOCKS))
    profile = corpus_profile(tmp_path)
    paths = ["--out", str(tmp_path / "trials.jsonl"), "--table", str(tmp_path / "table.json"),
             "--report", str(tmp_path / "report.md")]
    argv = [str(source), "--style", str(profile), "--models", "a,b", "--thinking", "low,high", "--groups", "1,4",
            "--baseline-model", "a", "--repeats", "1", "--words", "40", *paths]
    assert tune_generation.main(argv) == 0

    table = json.loads((tmp_path / "table.json").read_text())
    assert {(r["model"], r["thinking"], r["group_size"], r["measured"]) for r in table["rows"]} == {
        ("a", "low", 1, True), ("a", "low", 4, True), ("a", "high", 4, True), ("a", "high", 1, False),
        ("b", "low", 4, True), ("b", "high", 4, True), ("b", "low", 1, False), ("b", "high", 1, False)}
    assert all(r["ai"] == pytest.approx(0.4) and r["style"] is not None for r in table["rows"])
    assert table["default"]["group_size"] == 4 and "**Default" in (tmp_path / "report.md").read_text()
    labelled = [json.loads(line) for line in (tmp_path / "tuning_redteam.jsonl").read_text().splitlines()]
    assert len(labelled) == 5 and all(r["label"] == 1 for r in labelled)

    calls = len(gemini.calls)
    assert tune_generation.main(argv) == 0
    assert len(gemini.calls) == calls, "every trial is already saved"
    assert tune_generation.main(["--fit-only", "--daily-limit", "b=20", *paths]) == 0
    assert json.loads((tmp_path / "table.json").read_text())["daily_requests"] == {"b": 20}
    assert "b 20" in (tmp_path / "report.md").read_text()


def test_a_model_out_of_quota_is_skipped_for_the_rest_of_the_run_and_its_limit_kept(monkeypatch, tmp_path):
    def spent(*args, **kwargs):
        raise generate_essays.QuotaExhausted("a is out of quota: Rate limit exceeded for model a (limit: 1,500 "
                                             "requests per day on Free Tier). Please retry in 14h11m31s")

    monkeypatch.setattr(generate_essays, "chat", spent)
    monkeypatch.setattr(tune_generation, "load_env", lambda: None)
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    source = tmp_path / "notes.docx"
    source.write_bytes(document(BLOCKS))
    argv = [str(source), "--style", str(corpus_profile(tmp_path)), "--models", "a", "--repeats", "3",
            "--out", str(tmp_path / "trials.jsonl"), "--table", str(tmp_path / "table.json"),
            "--report", str(tmp_path / "report.md")]
    assert tune_generation.main(argv) == 1, "nothing was measured"
    assert not (tmp_path / "trials.jsonl").exists()
    assert json.loads((tmp_path / "trials_limits.json").read_text())["a"] == 1500
